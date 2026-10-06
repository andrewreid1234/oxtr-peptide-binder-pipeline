"""Group-meeting deck for the OXTR de novo cyclic peptide binder pipeline.

Ten slides. Built with python-pptx because this machine has no node/pptxgenjs.

Palette is the repo's figure palette (navy / figure-blue / figure-orange, red
reserved for "this fails"), so the embedded plots sit in the deck without
clashing — they were generated from the same colours.
"""
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as RGB
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

FIG = "/home/drewdog/projects/OXTR_peptides/docs/figures/slides/"
OUT = "/home/drewdog/.claude/jobs/07511682/tmp/deck/OXTR_pipeline_update.pptx"

NAVY, BLUE, ORANGE = RGB(0x16, 0x30, 0x5A), RGB(0x2E, 0x5F, 0xA3), RGB(0xD9, 0x7A, 0x29)
RED, GRAY, MUTED = RGB(0xB3, 0x36, 0x2B), RGB(0x8A, 0x8F, 0x98), RGB(0x6B, 0x70, 0x76)
INK, WHITE, ICE = RGB(0x22, 0x22, 0x22), RGB(0xFF, 0xFF, 0xFF), RGB(0xA8, 0xC0, 0xDE)
LIGHT = RGB(0xF2, 0xF5, 0xFA)
HEAD, BODY = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]


def tb(s, x, y, w, h, runs, size=16, color=INK, bold=False, italic=False,
       font=BODY, align=PP_ALIGN.LEFT, space_after=8, line=None, anchor=None):
    """runs: a string, or a list of (text, {overrides}) for mixed formatting."""
    box = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    items = [(runs, {})] if isinstance(runs, str) else runs
    for i, (text, o) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("space_after", space_after))
        if line or o.get("line"):
            p.line_spacing = o.get("line", line)
        if o.get("bullet"):
            p.level = 0
        r = p.add_run()
        r.text = ("•  " if o.get("bullet") else "") + text
        f = r.font
        f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
        f.italic = o.get("italic", italic); f.name = o.get("font", font)
        f.color.rgb = o.get("color", color)
    return box


def chip(s, x, y, w, h, text, fill, color=WHITE, size=14, bold=True):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = BODY
    r.font.color.rgb = color
    return sh


def dark(s):
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = NAVY


def footer(s, n):
    tb(s, 0.55, 7.02, 7, 0.3, "OXTR de novo cyclic peptide binders", size=10, color=GRAY)
    tb(s, 12.2, 7.02, 0.6, 0.3, str(n), size=10, color=GRAY, align=PP_ALIGN.RIGHT)


def title(s, text):
    tb(s, 0.55, 0.32, 12.2, 0.9, text, size=30, bold=True, color=INK, font=HEAD)


def takeaway(s, text):
    tb(s, 0.55, 6.38, 12.2, 0.66, text, size=14.5, color=MUTED, line=1.15)


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def figure(s, path, x, y, w, h):
    s.shapes.add_picture(path, I(x), I(y), I(w), I(h))


def stat(s, x, y, w, big, small, col):
    tb(s, x, y, w, 0.9, big, size=42, bold=True, color=col, font=HEAD,
       align=PP_ALIGN.CENTER, space_after=0)
    tb(s, x, y + 0.86, w, 0.95, small, size=12.5, color=MUTED,
       align=PP_ALIGN.CENTER, line=1.1)


# ------------------------------------------------------------------- 1
s = prs.slides.add_slide(BLANK); dark(s)
tb(s, 0.9, 2.15, 11.5, 1.8,
   "De novo cyclic peptide binders\nfor the oxytocin receptor",
   size=40, bold=True, color=WHITE, font=HEAD, line=1.08)
tb(s, 0.9, 4.15, 11.2, 1.5, [
   ("Scale-up complete through Stage 4, plus selectivity and permeability annotation", {"color": ICE}),
   ("265,700 designed   ·   143,595 docked   ·   3,000 scored with physics   ·   178 qualified",
    {"bold": True, "color": WHITE, "size": 18}),
], size=17, color=ICE)
notes(s, "Fifteen minutes. The story is that the funnel works, and that several things we were "
         "confident about turned out to be wrong when measured. Those reversals are the reason to "
         "believe the rest of it.")

# ------------------------------------------------------------------- 2
s = prs.slides.add_slide(BLANK); footer(s, 2)
title(s, "Oxytocin works in the brain. It cannot get there")
tb(s, 0.55, 1.4, 7.0, 4.6, [
   ("OXTR is a class A GPCR with a clear role in social behaviour, anxiety and stress, and an "
    "obvious starting ligand — oxytocin itself, a 9-residue disulfide-cyclised hormone with a "
    "solved receptor-bound structure.", {}),
   ("The obstacle is delivery. Oxytocin barely crosses the blood–brain barrier, so its central "
    "effects are hard to reach pharmacologically.", {}),
   ("So: design new disulfide-cyclised peptides from scratch against that structure, screen them "
    "computationally, and select a small set to synthesise.", {}),
   ("Binding is the hard constraint. Permeability is treated as something to engineer afterwards.",
    {"bold": True, "color": BLUE}),
], size=16, color=INK, line=1.25, space_after=13)
for i, (lab, col, tcol) in enumerate([
        ("Stage 1   RFdiffusion backbones", NAVY, WHITE),
        ("Stage 2   ProteinMPNN sequences", BLUE, WHITE),
        ("Stage 3   AfCycDesign docking", BLUE, WHITE),
        ("Stage 4   Rosetta binding energy", ORANGE, WHITE),
        ("Stage 7   Selectivity", GRAY, WHITE),
        ("Stage 8   Synthesis", LIGHT, NAVY)]):
    chip(s, 8.05, 1.5 + i * 0.72, 4.7, 0.56, lab, col, tcol, size=13.5)
notes(s, "Keep this short. Binding is the hard part; permeability we assumed we could engineer "
         "later. Slide eight shows that assumption is a bigger problem than we thought.")

# ------------------------------------------------------------------- 3
s = prs.slides.add_slide(BLANK); footer(s, 3)
title(s, "One real filter. The rest is how much compute we had")
figure(s, FIG + "slide_funnel.png", 1.12, 1.15, 11.08, 5.10)
takeaway(s, "Only the Stage 3 pocket gate is a pass/fail result — 87,338 of 143,595. Every other "
            "narrowing is a budget decision, and the 84,338 survivors we never scored were not rejected.")
notes(s, "This is the slide people misread. A funnel drawn without marking gates versus caps looks "
         "like five filters. There is one. We docked 54 percent of the pool and scored 3.4 percent "
         "of the survivors because of GPU and CPU time, not because the rest failed anything.")

# ------------------------------------------------------------------- 4
s = prs.slides.add_slide(BLANK); footer(s, 4)
title(s, "It worked: 7.6 REU better than chance")
figure(s, FIG + "slide_itworked.png", 1.47, 1.20, 10.38, 4.95)
takeaway(s, "3,000 candidates, each relaxed five times and averaged. The comparison is 200 randomly "
            "chosen Stage 3 survivors — the only honest 'no selection' baseline available.")
notes(s, "The baseline matters. Without it, minus 46 is a number with no meaning. Run-to-run noise "
         "was measured independently three times at about 3.5 REU, so the 7.6 gap is roughly twice "
         "the noise.")

# ------------------------------------------------------------------- 5
s = prs.slides.add_slide(BLANK); footer(s, 5)
title(s, "178 candidates clear everything we measured")
stat(s, 0.55, 1.35, 2.5, "178", "bind efficiently AND prefer OXTR\nover all three vasopressin receptors", BLUE)
stat(s, 3.15, 1.35, 2.2, "105", "distinct backbones —\nnot one scaffold", NAVY)
stat(s, 5.45, 1.35, 2.4, "53", "are 10 residues or shorter —\nthe size class we want", NAVY)
rows = [("Sequence", "Len", "dG/dSASA", "dG", "Select."),
        ("SGCLFGSCP", "9", "−3.682", "−51.1", "+0.168"),
        ("PRGCSFAFWPCD", "12", "−3.536", "−75.4", "+0.044"),
        ("GGPCGIFSFVCGP", "13", "−3.511", "−71.6", "+0.012"),
        ("SFCLGLDCPRA", "11", "−3.502", "−64.9", "+0.178"),
        ("GPGCFLGGDCP", "11", "−3.488", "−59.5", "+0.046"),
        ("GVCFGGGCDY", "10", "−3.480", "−44.4", "+0.045")]
tbl = s.shapes.add_table(len(rows), 5, I(8.15), I(1.35), I(4.65), I(2.5)).table
for w, c in zip((1.85, 0.5, 0.95, 0.7, 0.75), range(5)):
    tbl.columns[c].width = I(w)
for ri, row in enumerate(rows):
    tbl.rows[ri].height = I(0.33)
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.text = val
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = I(0.06); cell.margin_right = I(0.06)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
        f = p.runs[0].font
        f.size = Pt(11); f.name = BODY
        f.bold = ri == 0
        f.color.rgb = WHITE if ri == 0 else INK
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
tb(s, 0.55, 4.25, 7.3, 1.0,
   "Every one is a disulfide-cyclised macrocycle with a C-terminal amide — scored as the molecule "
   "we would actually make, not a linear approximation.", size=15, color=INK, line=1.2)
tb(s, 0.55, 5.35, 7.3, 1.0,
   "The −3.0 efficiency cut is a choice, not a derived threshold. It is the number on this slide "
   "I would defend least.", size=14, italic=True, color=RED, line=1.2)
notes(s, "178 against a need of about 100 — comfortably covered. Be honest about the threshold: "
         "minus 3.0 was chosen by eye, not derived. A second Rosetta batch of 3,000 is running now "
         "and should roughly double the pool.")

# ------------------------------------------------------------------- 6
s = prs.slides.add_slide(BLANK); footer(s, 6)
title(s, "You cannot pick a top 5. So make 100")
figure(s, FIG + "slide_shortlist.png", 1.61, 1.20, 10.12, 4.95)
takeaway(s, "Resampling the five structures per candidate: capturing the true best five with 90% "
            "confidence needs a shortlist of 86. A 20-candidate pilot had implied 8.")
notes(s, "This reframes the project. The leaders sit inside each other's error bars — only fourteen "
         "candidates are within two standard errors of the best. More sampling does not fix it, "
         "because the gap from rank one to rank five is about two standard errors to begin with. "
         "The answer is to synthesise breadth rather than rank harder.")

# ------------------------------------------------------------------- 7
s = prs.slides.add_slide(BLANK); footer(s, 7)
title(s, "Half the best binders prefer the wrong receptor")
figure(s, FIG + "slide_selectivity.png", 1.54, 1.20, 10.26, 4.95)
takeaway(s, "All 3,000 folded against AVPR1A, AVPR1B and AVPR2. AVPR2 — the renal antidiuretic "
            "receptor — is the worst offender, and selectivity is independent of binding rank (r = −0.065).")
notes(s, "Nothing upstream predicted this, which is why it was worth the GPU time. Read it "
         "asymmetrically: a negative margin is a red flag, but a positive margin is absence of "
         "evidence rather than proof of selectivity, because it is the same predictor.")

# ------------------------------------------------------------------- 8
s = prs.slides.add_slide(BLANK); footer(s, 8)
title(s, "Nothing we have made crosses passively")
figure(s, FIG + "slide_permeability.png", 0.47, 1.25, 12.40, 4.72)
takeaway(s, "Median TPSA 433 Å² against a passive-permeability ceiling near 140. And the BBB "
            "classifier cannot help: it scores leu-enkephalin, a confirmed non-permeant, at 0.959.")
notes(s, "Two findings. The physics: these are polar macrocycles and none is close to passive "
         "permeability. And the tool we would have used to judge that is unusable on this molecule "
         "class — proved by scoring three known non-permeants in the same batch. N-methylation is "
         "now the only computable lever we have.")

# ------------------------------------------------------------------- 9
s = prs.slides.add_slide(BLANK); footer(s, 9)
title(s, "A trap worth naming: you cannot grade your own selector")
figure(s, FIG + "slide_selector.png", 0.88, 1.25, 11.58, 4.75)
takeaway(s, "The feature we selected on shows no signal within the selected set — because that set "
            "holds only its top two values. Restricting the independent benchmark to the same range "
            "reproduces it exactly.")
notes(s, "I had this wrong for a day and wrote it up as a failed feature. It is range restriction. "
         "The honest check is against a sample spanning the feature's full range, which is what the "
         "independent 200 were. Worth saying out loud, because it will catch anyone who validates a "
         "selector on the set it selected.")

# ------------------------------------------------------------------ 10
s = prs.slides.add_slide(BLANK); dark(s)
tb(s, 0.9, 1.5, 11.5, 1.0, "Where this leaves us", size=38, bold=True, color=WHITE, font=HEAD)
tb(s, 0.9, 2.75, 11.4, 3.9, [
   ("178 candidates clear binding efficiency and selectivity — against a need for about 100.",
    {"bullet": True}),
   ("A second Rosetta batch of 3,000 is running and should roughly double that pool.",
    {"bullet": True}),
   ("Choose the synthesis set on diversity, not score order — the ordering inside the top 50 is not real.",
    {"bullet": True}),
   ("N-methylation scan is next, and the only permeability lever we can compute.", {"bullet": True}),
   ("None of this is binding data. The next milestone is wet-lab, and it is the only thing that "
    "settles any of it.", {"bullet": True, "bold": True, "color": WHITE}),
], size=16.5, color=ICE, line=1.2, space_after=14)
notes(s, "End on the honest note. Everything in this deck is predicted. The pipeline narrows a "
         "quarter of a million designs to a hundred defensibly, but it cannot tell you which one binds.")

prs.save(OUT)
print("wrote", OUT, "—", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
